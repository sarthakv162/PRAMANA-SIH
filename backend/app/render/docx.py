"""DOCX dossier renderer (§6.11) via python-docx. Evidence quote blocks are set in a
monospace font and never passed through any Markdown/HTML processing — the same "materialise
verbatim, don't reformat" rule §2 applies to every render target, not only the API response.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime

from docx import Document
from docx.shared import Pt

from app.render.dossier import DISCLAIMER, DossierItem


def render_docx(items: list[DossierItem], language: str) -> bytes:
    doc = Document()
    doc.add_heading("PRAMANA compliance dossier", level=0)
    meta_p = doc.add_paragraph()
    meta_p.add_run(f"Generated: {datetime.now(UTC).isoformat()} · Language: {language}").italic = True

    for i, item in enumerate(items, start=1):
        doc.add_heading(f"{i}. {item.title}", level=1)

        if item.kind == "missing":
            doc.add_paragraph().add_run(item.note or "").italic = True
            continue

        meta_bits = []
        if item.corpus_version:
            meta_bits.append(f"corpus version {item.corpus_version}")
        if item.as_of:
            meta_bits.append(f"as of {item.as_of}")
        if item.jurisdiction:
            meta_bits.append(f"jurisdiction {item.jurisdiction}")
        if meta_bits:
            doc.add_paragraph(" · ".join(meta_bits))

        for line in item.summary_lines:
            doc.add_paragraph(line)

        for quote in item.quotes:
            page = f", p.{quote.page}" if quote.page else ""
            label_p = doc.add_paragraph()
            label_p.add_run(f"{quote.citation_label}{page}").bold = True
            quote_p = doc.add_paragraph()
            quote_run = quote_p.add_run(quote.text)
            quote_run.font.name = "Courier New"
            quote_run.font.size = Pt(10)

        footer_p = doc.add_paragraph()
        footer_run = footer_p.add_run(
            f"Request: {item.request_id} · Receipt: {item.receipt_id} · "
            f"Entry hash: {item.entry_hash}"
        )
        footer_run.font.size = Pt(8)

    doc.add_paragraph()
    disclaimer_p = doc.add_paragraph()
    disclaimer_p.add_run(DISCLAIMER).italic = True

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
